"""
demo_manager_arquivos.py — Gera arquivos de demonstração para cada perfil.

Para cada perfil do demo_manager, produz arquivos que podem ser usados
em apresentações comerciais, videoaulas e testes manuais:

  demos/<perfil>/
    perfil.md                              ← descrição comercial/didática
    cartoes/fatura_<banco>_<mes>.xlsx     ← planilha importável (cartão)
    extratos/extrato_<banco>_<mes>.csv    ← extrato bancário Nubank-like
    extratos/extrato_<banco>_<mes>.ofx    ← extrato bancário OFX

Os arquivos refletem o histórico financeiro do perfil (gerado por
demo_manager.popular_modo_demo) — não são massa fake aleatória.

Uso:
    from demo_manager_arquivos import gerar_arquivos_perfil, gerar_todos
    gerar_arquivos_perfil("classe_media")   # 1 perfil
    gerar_todos()                           # os 20 perfis
"""

from __future__ import annotations

import csv
import logging
import tempfile
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

log = logging.getLogger(__name__)

_RAIZ = Path(__file__).parent / "demos"


# ─────────────────────────────────────────────────────────────────────────────
# Geração XLSX de fatura de cartão
# (mesmo formato lido por importar_fatura_modal — descricao, estabelecimento,
#  categoria, parcela, valor — copiado de exportar_model.gerar_template_*)
# ─────────────────────────────────────────────────────────────────────────────

_FILL_HEADER = PatternFill("solid", fgColor="1E3A5F")
_FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
_ALIGN_CTR   = Alignment(horizontal="center", vertical="center")


def _escrever_fatura_xlsx(linhas: list[dict], destino: Path) -> None:
    """Salva linhas no formato lido por importar_fatura_modal.

    Cada linha: {descricao, estabelecimento, categoria, parcela, valor}.
    """
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fatura"

    cols = ["descricao", "estabelecimento", "categoria", "parcela", "valor"]
    larguras = [32, 22, 18, 10, 12]
    for col_idx, (nome, larg) in enumerate(zip(cols, larguras), start=1):
        cell = ws.cell(row=1, column=col_idx, value=nome)
        cell.font      = _FONT_HEADER
        cell.fill      = _FILL_HEADER
        cell.alignment = _ALIGN_CTR
        ws.column_dimensions[chr(ord("A") + col_idx - 1)].width = larg

    for linha_idx, item in enumerate(linhas, start=2):
        ws.cell(row=linha_idx, column=1, value=item.get("descricao", ""))
        ws.cell(row=linha_idx, column=2, value=item.get("estabelecimento", ""))
        ws.cell(row=linha_idx, column=3, value=item.get("categoria", ""))
        ws.cell(row=linha_idx, column=4, value=item.get("parcela", ""))
        ws.cell(row=linha_idx, column=5, value=float(item.get("valor", 0.0)))

    wb.save(destino)
    log.info("Fatura XLSX salvo: %s (%d linhas)", destino.name, len(linhas))


# ─────────────────────────────────────────────────────────────────────────────
# Geração CSV Nubank de extrato
# (formato lido por extrato_parsers.nubank_csv:
#   Data,Valor,Identificador,Descrição)
# ─────────────────────────────────────────────────────────────────────────────

def _escrever_extrato_csv_nubank(
    linhas: list[dict], destino: Path
) -> None:
    """Salva extrato em CSV no formato Nubank.

    Cada linha: {data: 'YYYY-MM-DD', valor: float, descricao: str,
                 [identificador_unico: str]}
    """
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Data", "Valor", "Identificador", "Descrição"])
        for item in linhas:
            iso = item["data"]
            # Formato Nubank: DD/MM/AAAA
            ano, mes, dia = iso.split("-")
            data_br = f"{dia}/{mes}/{ano}"
            fitid = item.get("identificador_unico") or str(uuid.uuid4())
            valor = item["valor"]
            valor_str = f"{valor:.2f}"  # ponto decimal (CSV Nubank usa ponto)
            w.writerow([data_br, valor_str, fitid, item["descricao"]])

    log.info("Extrato CSV salvo: %s (%d linhas)", destino.name, len(linhas))


# ─────────────────────────────────────────────────────────────────────────────
# Geração OFX 1.0.2 SGML de extrato (formato Nubank)
# (lido por extrato_parsers.nubank_ofx)
# ─────────────────────────────────────────────────────────────────────────────

_OFX_HEADER = (
    "OFXHEADER:100\n"
    "DATA:OFXSGML\n"
    "VERSION:102\n"
    "SECURITY:NONE\n"
    "ENCODING:UTF-8\n"
    "CHARSET:NONE\n"
    "COMPRESSION:NONE\n"
    "OLDFILEUID:NONE\n"
    "NEWFILEUID:NONE\n"
)


def _escrever_extrato_ofx_nubank(
    linhas: list[dict], destino: Path,
    *,
    bankid: str = "0260", acct: str = "00000000-0",
) -> None:
    """Salva extrato em OFX 1.0.2 SGML (formato Nubank).

    Cada linha: {data: 'YYYY-MM-DD', valor: float, descricao: str,
                 [identificador_unico: str]}
    """
    destino.parent.mkdir(parents=True, exist_ok=True)
    if not linhas:
        log.warning("OFX vazio: %s", destino.name)
        return

    datas = sorted(item["data"] for item in linhas)
    dt_inicio = datas[0].replace("-", "") + "000000"
    dt_fim    = datas[-1].replace("-", "") + "000000"
    dt_server = datetime.now().strftime("%Y%m%d%H%M%S")

    bloco = [_OFX_HEADER, "<OFX>",
             "<SIGNONMSGSRSV1><SONRS>",
             "<STATUS><CODE>0</CODE><SEVERITY>INFO</SEVERITY></STATUS>",
             f"<DTSERVER>{dt_server}[0:GMT]</DTSERVER>",
             "<LANGUAGE>POR</LANGUAGE>",
             "<FI><ORG>NU PAGAMENTOS S.A.</ORG><FID>260</FID></FI>",
             "</SONRS></SIGNONMSGSRSV1>",
             "<BANKMSGSRSV1><STMTTRNRS>",
             "<TRNUID>1</TRNUID>",
             "<STATUS><CODE>0</CODE><SEVERITY>INFO</SEVERITY></STATUS>",
             "<STMTRS><CURDEF>BRL</CURDEF>",
             f"<BANKACCTFROM><BANKID>{bankid}</BANKID><BRANCHID>1</BRANCHID>"
             f"<ACCTID>{acct}</ACCTID><ACCTTYPE>CHECKING</ACCTTYPE></BANKACCTFROM>",
             "<BANKTRANLIST>",
             f"<DTSTART>{dt_inicio}[-3:BRT]</DTSTART>",
             f"<DTEND>{dt_fim}[-3:BRT]</DTEND>"]

    for item in linhas:
        valor = item["valor"]
        trntype = "CREDIT" if valor > 0 else "DEBIT"
        dtposted = item["data"].replace("-", "") + "000000"
        fitid = item.get("identificador_unico") or str(uuid.uuid4())
        memo = (item["descricao"] or "")[:255]
        bloco.append("<STMTTRN>")
        bloco.append(f"<TRNTYPE>{trntype}</TRNTYPE>")
        bloco.append(f"<DTPOSTED>{dtposted}[-3:BRT]</DTPOSTED>")
        bloco.append(f"<TRNAMT>{valor:.2f}</TRNAMT>")
        bloco.append(f"<FITID>{fitid}</FITID>")
        bloco.append(f"<MEMO>{memo}</MEMO>")
        bloco.append("</STMTTRN>")

    bloco.extend([
        "</BANKTRANLIST>",
        "</STMTRS></STMTTRNRS></BANKMSGSRSV1>",
        "</OFX>",
    ])

    destino.write_text("\n".join(bloco), encoding="utf-8")
    log.info("Extrato OFX salvo: %s (%d linhas)", destino.name, len(linhas))


# ─────────────────────────────────────────────────────────────────────────────
# Coleta de dados do perfil → dicts pra geração de arquivos
# ─────────────────────────────────────────────────────────────────────────────

def _coletar_fatura_perfil(conn, cartao_id: int, mes: int, ano: int) -> list[dict]:
    """Pega parcelas pendentes/pagas do cartão no mês como linhas de fatura."""
    ref = f"{ano:04d}-{mes:02d}"
    rows = conn.execute(
        """
        SELECT cc.descricao, cc.estabelecimento, cc.categoria,
               pc.numero_parcela, pc.valor, cc.total_parcelas
          FROM parcelas_cartao pc
          JOIN compras_cartao cc ON cc.id = pc.compra_id
         WHERE pc.cartao_id = ? AND pc.mes_referencia = ?
         ORDER BY cc.descricao
        """,
        (cartao_id, ref),
    ).fetchall()
    return [{
        "descricao":       r["descricao"] or "",
        "estabelecimento": r["estabelecimento"] or "",
        "categoria":       r["categoria"] or "",
        "parcela":         f"{r['numero_parcela']}/{r['total_parcelas']}",
        "valor":           float(r["valor"] or 0.0),
    } for r in rows]


def _coletar_extrato_perfil(conn, mes: int, ano: int) -> list[dict]:
    """Constrói linhas de extrato fictício a partir do que o perfil tem
    no mês: contas_pagar (saídas) + receitas (entradas).

    Como o demo NÃO popula lancamentos_banco (esses são importados pelo
    user), geramos linhas espelhando contas_pagar + fontes_receita pra ter
    um arquivo demonstrativo coerente com o cenário.
    """
    inicio = f"{ano:04d}-{mes:02d}-01"
    fim = (f"{ano + 1:04d}-01-01" if mes == 12
           else f"{ano:04d}-{mes + 1:02d}-01")
    linhas: list[dict] = []

    # Débitos: contas_pagar do mês
    rows = conn.execute(
        """
        SELECT cp.data_vencimento, cp.descricao, cp.valor,
               COALESCE(pc.nome, 'Despesa') AS cat
          FROM contas_pagar cp
          LEFT JOIN plano_contas pc ON pc.id = cp.plano_conta_id
         WHERE cp.data_vencimento >= ? AND cp.data_vencimento < ?
           AND cp.status != 'cancelado'
         ORDER BY cp.data_vencimento
        """,
        (inicio, fim),
    ).fetchall()
    for r in rows:
        desc = r["descricao"] or r["cat"]
        linhas.append({
            "data":      r["data_vencimento"],
            "valor":     -abs(float(r["valor"] or 0.0)),  # débito
            "descricao": desc,
        })

    # Créditos: fontes_receita ativas — 1 entrada no dia 5 do mês
    import calendar
    dia_pag = min(5, calendar.monthrange(ano, mes)[1])
    rows_r = conn.execute(
        "SELECT nome, valor_mensal FROM fontes_receita WHERE ativa = 1"
    ).fetchall()
    for r in rows_r:
        linhas.append({
            "data":      f"{ano:04d}-{mes:02d}-{dia_pag:02d}",
            "valor":     float(r["valor_mensal"] or 0.0),
            "descricao": r["nome"] or "Receita",
        })

    linhas.sort(key=lambda x: x["data"])
    return linhas


# ─────────────────────────────────────────────────────────────────────────────
# Entry points públicos
# ─────────────────────────────────────────────────────────────────────────────

def gerar_arquivos_perfil(
    perfil: str,
    *,
    raiz: Optional[Path] = None,
    meses: int = 3,
) -> dict[str, int]:
    """Gera arquivos demo (XLSX/CSV/OFX) pro perfil em demos/<perfil>/.

    Cria um banco temporário, popula com o perfil e extrai os dados pra
    arquivos. NÃO toca no banco real do usuário.

    Retorna estatísticas: {'faturas': N, 'extratos_csv': N, 'extratos_ofx': N}
    """
    import database as db
    from demo_manager import popular_modo_demo, PERFIS_DEMO

    if perfil not in PERFIS_DEMO:
        raise ValueError(f"Perfil desconhecido: {perfil!r}")

    raiz = raiz or _RAIZ
    pasta_perfil = raiz / perfil
    pasta_perfil.mkdir(parents=True, exist_ok=True)

    # README do perfil
    md = pasta_perfil / "perfil.md"
    md.write_text(
        f"# {PERFIS_DEMO[perfil]}\n\n"
        f"Chave: `{perfil}`\n\n"
        f"Arquivos gerados automaticamente por `demo_manager_arquivos`.\n"
        f"Para popular o sistema com este perfil:\n"
        f"```python\n"
        f"from demo_manager import popular_modo_demo\n"
        f"popular_modo_demo({perfil!r})\n"
        f"```\n\n"
        f"## Arquivos\n\n"
        f"- `cartoes/` — faturas XLSX importáveis (1 por cartão x mês)\n"
        f"- `extratos/` — extratos bancários CSV e OFX (Nubank-like)\n",
        encoding="utf-8",
    )

    # Setup banco temporário pra extrair dados do perfil
    tmp = tempfile.mkdtemp(prefix=f"serenus_demo_{perfil}_")
    db_path = Path(tmp) / "demo.db"
    orig_caminho = db.CAMINHO_BANCO
    db.CAMINHO_BANCO = db_path
    try:
        db.inicializar_banco()
        popular_modo_demo(perfil)

        hoje = date.today()
        stats = {"faturas": 0, "extratos_csv": 0, "extratos_ofx": 0}

        with db.conectar() as conn:
            # Faturas: 1 arquivo por cartão x mes_referencia, pros últimos N meses
            cartoes = conn.execute("SELECT id, nome FROM cartoes WHERE ativo = 1").fetchall()
            for off in range(-(meses - 1), 1):
                mes = ((hoje.month - 1 + off) % 12) + 1
                ano = hoje.year + (hoje.month - 1 + off) // 12
                for c in cartoes:
                    linhas = _coletar_fatura_perfil(conn, c["id"], mes, ano)
                    if not linhas:
                        continue
                    nome_slug = "".join(ch if ch.isalnum() else "_"
                                          for ch in c["nome"].lower())
                    destino = pasta_perfil / "cartoes" / \
                              f"fatura_{nome_slug}_{ano:04d}-{mes:02d}.xlsx"
                    _escrever_fatura_xlsx(linhas, destino)
                    stats["faturas"] += 1

            # Extrato consolidado: 1 CSV + 1 OFX por mês recente
            for off in range(-(meses - 1), 1):
                mes = ((hoje.month - 1 + off) % 12) + 1
                ano = hoje.year + (hoje.month - 1 + off) // 12
                linhas = _coletar_extrato_perfil(conn, mes, ano)
                if not linhas:
                    continue
                csv_dest = pasta_perfil / "extratos" / \
                           f"extrato_nubank_{ano:04d}-{mes:02d}.csv"
                ofx_dest = pasta_perfil / "extratos" / \
                           f"extrato_nubank_{ano:04d}-{mes:02d}.ofx"
                _escrever_extrato_csv_nubank(linhas, csv_dest)
                _escrever_extrato_ofx_nubank(linhas, ofx_dest)
                stats["extratos_csv"] += 1
                stats["extratos_ofx"] += 1
    finally:
        db.CAMINHO_BANCO = orig_caminho

    return stats


def gerar_todos(*, raiz: Optional[Path] = None, meses: int = 3) -> dict[str, dict]:
    """Gera arquivos pra TODOS os perfis. Retorna {perfil: stats}."""
    from demo_manager import PERFIS_DEMO
    resultado = {}
    for perfil in PERFIS_DEMO:
        log.info("Gerando arquivos do perfil: %s", perfil)
        resultado[perfil] = gerar_arquivos_perfil(perfil, raiz=raiz, meses=meses)
    return resultado


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    if len(sys.argv) > 1:
        perfil = sys.argv[1]
        stats = gerar_arquivos_perfil(perfil)
        print(f"OK {perfil}: {stats}")
    else:
        stats_all = gerar_todos()
        print("\nResumo:")
        for p, s in stats_all.items():
            print(f"  {p:<26s} {s}")
