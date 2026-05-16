"""
test_exportar_model.py — Testes do módulo de exportação CSV/Excel.
"""
import pytest
from pathlib import Path
from datetime import datetime


# ── helpers de dados ─────────────────────────────────────────────────────────

def _inserir_plano(nome="Aluguel"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO plano_contas (nome,tipo_custo,categoria,ativa,padrao,criado_em)"
            " VALUES (?,'fixo','Moradia',1,0,?)", (nome, agora))
        return cur.lastrowid


def _inserir_conta_pagar(plano_id, valor, venc, status="pago"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        conn.execute(
            "INSERT INTO contas_pagar (plano_conta_id,descricao,valor,"
            "data_vencimento,status,recorrente,criado_em) VALUES (?,?,?,?,?,0,?)",
            (plano_id, "Desc", valor, venc, status, agora))


def _inserir_fonte_receita(valor=3000.0):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        conn.execute(
            "INSERT INTO fontes_receita (nome,valor_mensal,periodicidade,"
            "dia_pagamento,ativa,criado_em) VALUES ('Salário',?,?,5,1,?)",
            (valor, "mensal", agora))


def _inserir_conta_inv():
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO contas_investimento (nome,instituicao,tipo,ativa,criado_em)"
            " VALUES ('Corretora','Banco','corretora',1,?)", (agora,))
        return cur.lastrowid


def _inserir_ativo(conta_id, codigo="TEST4", tipo="acao"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO ativos (codigo,nome,tipo,conta_investimento_id,ativo,criado_em)"
            " VALUES (?,?,?,?,1,?)", (codigo, "Empresa Teste", tipo, conta_id, agora))
        return cur.lastrowid


def _inserir_posicao(ativo_id, qtd, custo, valor_inv, valor_atual, lucro_real):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO posicao_cache"
            " (ativo_id,quantidade_atual,custo_medio,valor_investido,"
            "valor_atual,lucro_realizado,atualizado_em) VALUES (?,?,?,?,?,?,?)",
            (ativo_id, qtd, custo, valor_inv, valor_atual, lucro_real, agora))


def _inserir_cartao():
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO cartoes (nome,banco,bandeira,limite,limite_disponivel,ativo,criado_em)"
            " VALUES ('Nu Teste','nubank','visa',5000,3000,1,?)", (agora,))
        return cur.lastrowid


def _inserir_parcela(cartao_id, mes_ref, valor=500.0):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO compras_cartao (cartao_id,descricao,valor_total,valor_parcela,"
            "total_parcelas,parcelas_pagas,mes_inicio,categoria,estabelecimento,criado_em)"
            " VALUES (?,?,?,?,1,0,?,?,?,?)",
            (cartao_id, "Compra X", valor, valor, mes_ref, "", "", agora))
        conn.execute(
            "INSERT INTO parcelas_cartao (compra_id,cartao_id,numero_parcela,"
            "mes_referencia,valor,status) VALUES (?,?,1,?,?,'pendente')",
            (cur.lastrowid, cartao_id, mes_ref, valor))


# ─────────────────────────────────────────────────────────────────────────────
# Exportação de extrato
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarExtrato:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_extrato_xlsx
        plano = _inserir_plano()
        _inserir_conta_pagar(plano, 1200.0, "2025-06-05")
        _inserir_fonte_receita()

        dest = tmp_path / "extrato.xlsx"
        exportar_extrato_xlsx(mes=6, ano=2025, caminho=str(dest))
        assert dest.exists()
        assert dest.stat().st_size > 0

    def test_xlsx_tem_linhas_de_dados(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_extrato_xlsx
        plano = _inserir_plano()
        _inserir_conta_pagar(plano, 800.0, "2025-06-10")
        _inserir_fonte_receita(2500.0)

        dest = tmp_path / "extrato.xlsx"
        exportar_extrato_xlsx(mes=6, ano=2025, caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        # linha 1 = cabeçalho, deve ter pelo menos 2 linhas (cabeçalho + dado)
        assert ws.max_row >= 2

    def test_xlsx_cabecalho_correto(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_extrato_xlsx

        dest = tmp_path / "extrato.xlsx"
        exportar_extrato_xlsx(mes=6, ano=2025, caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Data"     in headers
        assert "Descrição" in headers
        assert "Débito"   in headers
        assert "Crédito"  in headers


# ─────────────────────────────────────────────────────────────────────────────
# Exportação de carteira
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarCarteira:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_carteira_xlsx
        ci = _inserir_conta_inv()
        ativo = _inserir_ativo(ci)
        _inserir_posicao(ativo, 100, 30.0, 3000.0, 3500.0, 0.0)

        dest = tmp_path / "carteira.xlsx"
        exportar_carteira_xlsx(caminho=str(dest))
        assert dest.exists()
        assert dest.stat().st_size > 0

    def test_cabecalho_carteira(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_carteira_xlsx

        dest = tmp_path / "carteira.xlsx"
        exportar_carteira_xlsx(caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Código"         in headers
        assert "Valor Investido" in headers
        assert "Valor Atual"    in headers

    def test_carteira_vazia_gera_apenas_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_carteira_xlsx

        dest = tmp_path / "carteira.xlsx"
        exportar_carteira_xlsx(caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        assert ws.max_row == 1  # só cabeçalho


# ─────────────────────────────────────────────────────────────────────────────
# Exportação de fatura
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarFatura:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_fatura_xlsx
        cartao = _inserir_cartao()
        _inserir_parcela(cartao, "2025-06")

        dest = tmp_path / "fatura.xlsx"
        exportar_fatura_xlsx(cartao_id=cartao, mes=6, ano=2025, caminho=str(dest))
        assert dest.exists()
        assert dest.stat().st_size > 0

    def test_fatura_tem_parcelas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_fatura_xlsx
        cartao = _inserir_cartao()
        _inserir_parcela(cartao, "2025-07", 300.0)
        _inserir_parcela(cartao, "2025-07", 200.0)

        dest = tmp_path / "fatura.xlsx"
        exportar_fatura_xlsx(cartao_id=cartao, mes=7, ano=2025, caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        assert ws.max_row >= 3  # cabeçalho + 2 parcelas

    def test_cabecalho_fatura(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_fatura_xlsx
        cartao = _inserir_cartao()

        dest = tmp_path / "fatura.xlsx"
        exportar_fatura_xlsx(cartao_id=cartao, mes=6, ano=2025, caminho=str(dest))

        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        # linha 1 = título merged; linha 2 = cabeçalho real
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        assert "Descrição" in headers
        assert "Valor"     in headers
        assert "Status"    in headers
