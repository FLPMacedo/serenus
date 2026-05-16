"""
test_exportar_todos.py — Testes dos relatórios Excel de todos os módulos.

Módulos cobertos:
  - Contas a Pagar  → exportar_contas_pagar_xlsx
  - Dívidas         → exportar_dividas_xlsx
  - Receitas        → exportar_receitas_xlsx
  - Projeção        → exportar_projecao_xlsx
  - Metas           → exportar_metas_xlsx
  - IR Estimado     → exportar_ir_xlsx
"""
import pytest
from datetime import datetime, date


# ─────────────────────────────────────────────────────────────────────────────
# Helpers de inserção
# ─────────────────────────────────────────────────────────────────────────────

def _plano(nome="Aluguel", tipo="fixo"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO plano_contas (nome,tipo_custo,categoria,ativa,padrao,criado_em)"
            " VALUES (?,?,'Moradia',1,0,?)", (nome, tipo, agora))
        return cur.lastrowid


def _conta_pagar(plano_id, valor, venc, status="pendente", desc="Desc"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        conn.execute(
            "INSERT INTO contas_pagar (plano_conta_id,descricao,valor,"
            "data_vencimento,status,recorrente,criado_em) VALUES (?,?,?,?,?,0,?)",
            (plano_id, desc, valor, venc, status, agora))


def _divida(nome="Financiamento", saldo=10000.0, parcela=500.0,
            total_p=24, pagas=2, taxa=1.5):
    from views.visao_longo_prazo.divida_model import salvar_divida
    return salvar_divida({
        "nome": nome, "tipo": "financiamento",
        "saldo_atual": saldo, "parcela_mensal": parcela,
        "total_parcelas": total_p, "parcelas_pagas": pagas,
        "dia_vencimento": 10, "taxa_juros": taxa, "observacao": "",
    })


def _fonte(nome="Salário", valor=5000.0, tipo="clt", period="mensal"):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO fontes_receita (nome,tipo,valor_mensal,periodicidade,"
            "dia_pagamento,ativa,criado_em) VALUES (?,?,?,?,5,1,?)",
            (nome, tipo, valor, period, agora))
        return cur.lastrowid


def _limpar_dividas():
    from database import conectar
    with conectar() as conn:
        conn.execute("DELETE FROM dividas")


def _limpar_fontes():
    from database import conectar
    with conectar() as conn:
        conn.execute("DELETE FROM fontes_receita")
        conn.execute("DELETE FROM receitas_especiais")


def _especial(nome="13º", mes=11, valor=5000.0):
    from views.receitas.receita_model import salvar_receita_especial
    return salvar_receita_especial({
        "nome": nome, "mes": mes, "valor": valor,
        "tipo": "clt", "recorrente_anual": True,
    })


def _meta(nome="Viagem", alvo=10000.0, atual=2000.0, prazo=None):
    from views.metas.metas_model import salvar_meta
    return salvar_meta({
        "nome": nome, "valor_alvo": alvo, "valor_atual": atual,
        "prazo": prazo, "descricao": "",
    })


def _ativo_com_lucro(lucro=1000.0):
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    from database import conectar
    with conectar() as conn:
        cur = conn.execute(
            "INSERT INTO contas_investimento (nome,instituicao,tipo,ativa,criado_em)"
            " VALUES ('C','B','corretora',1,?)", (agora,))
        cid = cur.lastrowid
        cur = conn.execute(
            "INSERT INTO ativos (codigo,nome,tipo,conta_investimento_id,ativo,criado_em)"
            " VALUES ('TST4','Teste',?,?,1,?)", ("acao", cid, agora))
        aid = cur.lastrowid
        conn.execute(
            "INSERT OR REPLACE INTO posicao_cache"
            " (ativo_id,quantidade_atual,custo_medio,valor_investido,"
            "valor_atual,lucro_realizado,atualizado_em) VALUES (?,10,10,100,200,?,?)",
            (aid, lucro, agora))
        return aid


# ─────────────────────────────────────────────────────────────────────────────
# Contas a Pagar
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarContasPagar:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        pid = _plano()
        _conta_pagar(pid, 1200.0, "2025-06-10")
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nomes_abas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        assert "Despesas 06-2025" in wb.sheetnames
        assert "Resumo" in wb.sheetnames

    def test_cabecalho_despesas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Despesas 06-2025"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Conta" in headers
        assert "Valor" in headers
        assert "Vencimento" in headers
        assert "Status" in headers

    def test_linhas_com_dados(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        pid = _plano()
        _conta_pagar(pid, 800.0, "2025-06-05")
        _conta_pagar(pid, 300.0, "2025-06-15")
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Despesas 06-2025"]
        assert ws.max_row >= 3  # cabeçalho + 2 linhas

    def test_vazio_gera_apenas_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Despesas 06-2025"]
        assert ws.max_row == 1

    def test_total_linha_presente(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        pid = _plano()
        _conta_pagar(pid, 1000.0, "2025-06-10")
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Despesas 06-2025"]
        # última linha deve ter "TOTAL"
        ultima = ws.max_row
        vals = [ws.cell(ultima, c).value for c in range(1, 3)]
        assert "TOTAL" in vals

    def test_aba_resumo_tem_status(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        pid = _plano()
        _conta_pagar(pid, 500.0, "2025-06-10", status="pago")
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Resumo"]
        valores = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any(v in ("Pendente", "Pago", "Status") for v in valores if v)

    def test_periodo_correto_filtra_mes(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_contas_pagar_xlsx
        pid = _plano()
        _conta_pagar(pid, 400.0, "2025-06-10")   # mês 6
        _conta_pagar(pid, 900.0, "2025-07-10")   # mês 7 — não deve aparecer
        dest = tmp_path / "cp.xlsx"
        exportar_contas_pagar_xlsx(6, 2025, str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Despesas 06-2025"]
        assert ws.max_row == 3  # cabeçalho + 1 dado + total


# ─────────────────────────────────────────────────────────────────────────────
# Dívidas
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarDividas:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_dividas_xlsx
        _divida()
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nomes_abas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert "Dívidas Ativas" in wb.sheetnames
        assert "Projeção" in wb.sheetnames

    def test_cabecalho_dividas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Dívidas Ativas"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Nome" in headers
        assert "Saldo Atual" in headers
        assert "Parcela/Mês" in headers
        assert "Quitação" in headers

    def test_linhas_com_dados(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        _divida("Empréstimo A")
        _divida("Cartão B", saldo=5000.0, parcela=200.0, total_p=36, pagas=5)
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Dívidas Ativas"]
        assert ws.max_row >= 3  # cabeçalho + 2

    def test_vazio_gera_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        _limpar_dividas()
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Dívidas Ativas"]
        assert ws.max_row == 1

    def test_total_saldo_presente(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        _divida(saldo=8000.0)
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Dívidas Ativas"]
        ultima = ws.max_row
        vals = [ws.cell(ultima, c).value for c in range(1, 3)]
        assert "TOTAL" in vals

    def test_projecao_tem_cabecalho_correto(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Projeção"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Mês" in headers
        assert "Saldo Total" in headers

    def test_projecao_vazia_tem_so_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_dividas_xlsx
        _limpar_dividas()
        dest = tmp_path / "div.xlsx"
        exportar_dividas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Projeção"]
        assert ws.max_row == 1  # sem dívidas, projeção vazia


# ─────────────────────────────────────────────────────────────────────────────
# Receitas
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarReceitas:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_receitas_xlsx
        _fonte()
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nomes_abas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert "Fontes de Receita" in wb.sheetnames
        assert "Receitas Especiais" in wb.sheetnames

    def test_cabecalho_fontes(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Fontes de Receita"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Nome" in headers
        assert "Valor Mensal" in headers
        assert "Periodicidade" in headers

    def test_cabecalho_especiais(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Receitas Especiais"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Nome" in headers
        assert "Mês" in headers
        assert "Valor" in headers
        assert "Recorrente" in headers

    def test_linhas_fontes(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        _limpar_fontes()
        _fonte("Salário", 5000.0)
        _fonte("Freela", 1500.0, tipo="freela")
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Fontes de Receita"]
        assert ws.max_row >= 3  # cabeçalho + 2 + total

    def test_linhas_especiais(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        _especial("13º", mes=11, valor=5000.0)
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Receitas Especiais"]
        assert ws.max_row >= 2  # cabeçalho + 1 dado

    def test_vazio_gera_cabecalhos(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        _limpar_fontes()
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert wb["Fontes de Receita"].max_row == 1
        assert wb["Receitas Especiais"].max_row == 1

    def test_total_fontes_presente(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_receitas_xlsx
        _fonte(valor=3000.0)
        dest = tmp_path / "rec.xlsx"
        exportar_receitas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Fontes de Receita"]
        ultima = ws.max_row
        vals = [ws.cell(ultima, c).value for c in range(1, 3)]
        assert "TOTAL" in vals


# ─────────────────────────────────────────────────────────────────────────────
# Projeção
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarProjecao:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nome_aba_padrao(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert any("Projeção" in s for s in wb.sheetnames)

    def test_cabecalho_projecao(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Mês" in headers
        assert "Receitas" in headers
        assert "Total Saídas" in headers
        assert "Saldo" in headers

    def test_tem_60_linhas_default(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        assert ws.max_row == 61  # 1 cabeçalho + 60 meses

    def test_parametro_meses_personalizado(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest), meses=12)
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        assert ws.max_row == 13  # 1 cabeçalho + 12 meses

    def test_coluna_mes_preenchida(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_projecao_xlsx
        dest = tmp_path / "proj.xlsx"
        exportar_projecao_xlsx(str(dest), meses=3)
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        # célula A2 deve ser um label de mês (não vazio)
        assert ws.cell(2, 1).value is not None


# ─────────────────────────────────────────────────────────────────────────────
# Metas
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarMetas:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_metas_xlsx
        _meta()
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nome_aba(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert "Metas" in wb.sheetnames

    def test_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Metas"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Nome" in headers
        assert "Valor Alvo" in headers
        assert "Valor Atual" in headers
        assert "Progresso %" in headers

    def test_linhas_com_dados(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        _meta("Viagem", 10000.0, 3000.0)
        _meta("Carro",  50000.0, 10000.0)
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Metas"]
        assert ws.max_row >= 3  # cabeçalho + 2

    def test_vazio_gera_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Metas"]
        assert ws.max_row == 1

    def test_total_alvo_e_atual(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        _meta("A", 5000.0, 1000.0)
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Metas"]
        ultima = ws.max_row
        vals = [ws.cell(ultima, c).value for c in range(1, 3)]
        assert "TOTAL" in vals

    def test_meta_com_prazo_exporta_data(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_metas_xlsx
        _meta("Com Prazo", 1000.0, 0.0, prazo="2027-12-31")
        dest = tmp_path / "metas.xlsx"
        exportar_metas_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["Metas"]
        # Linha 2 (primeira meta), coluna 5 = Prazo
        prazo_val = ws.cell(2, 5).value
        assert prazo_val is not None and prazo_val != "—"


# ─────────────────────────────────────────────────────────────────────────────
# IR Estimado
# ─────────────────────────────────────────────────────────────────────────────

class TestExportarIR:
    def test_gera_arquivo_xlsx(self, banco, tmp_path):
        from views.exportar.exportar_model import exportar_ir_xlsx
        _ativo_com_lucro()
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        assert dest.exists() and dest.stat().st_size > 0

    def test_nome_aba(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        assert "IR Estimado" in wb.sheetnames

    def test_cabecalho(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["IR Estimado"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Código" in headers
        assert "Lucro Realizado" in headers
        assert "Alíquota %" in headers
        assert "IR Estimado" in headers

    def test_vazio_sem_lucro(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["IR Estimado"]
        assert ws.max_row == 1  # apenas cabeçalho

    def test_com_lucro_tem_linhas(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        _ativo_com_lucro(500.0)
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["IR Estimado"]
        assert ws.max_row >= 2

    def test_total_ir_presente(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        _ativo_com_lucro(1000.0)
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["IR Estimado"]
        ultima = ws.max_row
        # última linha ou penúltima (pode ter nota de rodapé)
        found_total = False
        for r in range(2, ultima + 1):
            if ws.cell(r, 1).value == "TOTAL":
                found_total = True
                break
        assert found_total

    def test_nota_rodape_presente(self, banco, tmp_path):
        import openpyxl
        from views.exportar.exportar_model import exportar_ir_xlsx
        _ativo_com_lucro(200.0)
        dest = tmp_path / "ir.xlsx"
        exportar_ir_xlsx(str(dest))
        wb = openpyxl.load_workbook(dest)
        ws = wb["IR Estimado"]
        # deve haver alguma célula com texto de aviso
        all_vals = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        assert any(v and "Estimativa" in str(v) for v in all_vals)
