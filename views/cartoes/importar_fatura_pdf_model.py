"""
importar_fatura_pdf_model.py — Pipeline de importação de fatura via PDF.

Etapa 1: classes de erro, detecção de senha, extração de texto base.
Etapas seguintes: parsers por layout, pipeline completo, OCR fallback.

Usa pypdf (já presente nas dependências) tanto para detecção de senha
quanto para descriptografia e extração de texto. pdfplumber pode ser
adicionado depois para extração de tabelas em parsers específicos.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Erros explícitos do módulo
# ---------------------------------------------------------------------------

class PDFSenhaIncorretaError(Exception):
    """PDF protegido por senha e a senha fornecida não funcionou (ou ausente)."""


class PDFCorrompidoError(Exception):
    """Arquivo inexistente, ilegível, sem header PDF válido ou ilegível pelo pypdf."""


class PDFLayoutDesconhecidoError(Exception):
    """Nenhum parser concreto reconheceu o layout do PDF (usado em etapas 2+)."""


class PDFCamposIncompletosError(Exception):
    """Parser identificou o layout mas faltam campos obrigatórios (usado em etapas 3+)."""


# ---------------------------------------------------------------------------
# Detecção de senha
# ---------------------------------------------------------------------------

def pdf_tem_senha(caminho: str) -> bool:
    """Retorna True se o PDF estiver protegido por senha.

    Levanta PDFCorrompidoError se o arquivo não puder ser aberto como PDF.
    """
    p = Path(caminho)
    if not p.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:  # pragma: no cover — pypdf é dependência base
        raise RuntimeError(
            "Biblioteca pypdf não está disponível. Reinstale as dependências."
        ) from e

    try:
        reader = PdfReader(str(p))
        tem = bool(reader.is_encrypted)
        log.info("pdf_tem_senha(%s) = %s", p.name, tem)
        return tem
    except PdfReadError as e:
        log.warning("pdf_tem_senha: PDF corrompido %s: %s", p.name, e)
        raise PDFCorrompidoError(
            f"Arquivo {p.name!r} não é um PDF válido: {e}"
        ) from e
    except Exception as e:
        log.warning("pdf_tem_senha: erro inesperado em %s: %s", p.name, e)
        raise PDFCorrompidoError(
            f"Falha ao abrir {p.name!r}: {e}"
        ) from e


# ---------------------------------------------------------------------------
# Extração de texto
# ---------------------------------------------------------------------------

def extrair_texto_pdf(caminho: str, senha: Optional[str] = None) -> str:
    """Extrai todo o texto do PDF.

    - Se o PDF estiver protegido e a senha for None ou incorreta, levanta
      PDFSenhaIncorretaError.
    - Se a senha for fornecida para um PDF não protegido, é ignorada.
    - Levanta PDFCorrompidoError se o arquivo for inválido.

    Retorna string com o texto concatenado de todas as páginas.
    """
    p = Path(caminho)
    if not p.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(p))
    except PdfReadError as e:
        raise PDFCorrompidoError(
            f"Arquivo {p.name!r} não é um PDF válido: {e}"
        ) from e
    except Exception as e:
        raise PDFCorrompidoError(
            f"Falha ao abrir {p.name!r}: {e}"
        ) from e

    if reader.is_encrypted:
        if not senha:
            log.info("extrair_texto_pdf(%s): protegido, senha ausente", p.name)
            raise PDFSenhaIncorretaError(
                f"{p.name} é protegido por senha. Forneça a senha."
            )
        try:
            # pypdf.decrypt retorna PasswordType (0 = falhou, 1 = user, 2 = owner)
            resultado = reader.decrypt(senha)
            if not resultado:
                raise PDFSenhaIncorretaError(
                    f"Senha incorreta para {p.name}."
                )
        except PDFSenhaIncorretaError:
            raise
        except Exception as e:
            log.warning("extrair_texto_pdf: erro descriptografando %s: %s", p.name, e)
            raise PDFSenhaIncorretaError(
                f"Falha ao descriptografar {p.name}: {e}"
            ) from e

    try:
        partes: list[str] = []
        for pag in reader.pages:
            try:
                partes.append(pag.extract_text() or "")
            except Exception as e:  # pragma: no cover — falha por página
                log.warning("extract_text de página falhou em %s: %s", p.name, e)
        texto = "\n".join(partes).strip()
    except Exception as e:
        raise PDFCorrompidoError(
            f"Falha ao extrair texto de {p.name!r}: {e}"
        ) from e

    if not texto:
        log.warning("extrair_texto_pdf(%s): texto vazio (PDF escaneado?)", p.name)
    else:
        log.info("extrair_texto_pdf(%s): %d caracteres", p.name, len(texto))
    return texto


# ---------------------------------------------------------------------------
# Pipeline completo
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# OCR fallback (graceful degradation)
# ---------------------------------------------------------------------------

# Limite abaixo do qual consideramos que o pypdf não conseguiu extrair texto
# útil — provavelmente PDF puramente imagem (escaneado).
_LIMITE_TEXTO_VAZIO = 50


def ocr_disponivel() -> bool:
    """True se as bibliotecas pytesseract + pdf2image estiverem instaladas
    e Tesseract for executável. Usado para escolher entre tentar OCR ou
    devolver mensagem clara ao usuário."""
    try:
        import pytesseract  # noqa: F401
        import pdf2image    # noqa: F401
    except ImportError:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _extrair_texto_via_ocr(caminho: str, senha: Optional[str] = None) -> str:
    """Faz OCR de cada página do PDF e retorna texto concatenado.

    Pré-condição: ocr_disponivel() == True. Senha (se houver) é usada pra
    descriptografar antes de converter páginas em imagem.
    """
    import pdf2image
    import pytesseract

    log.info("OCR: convertendo páginas em imagem para %s", caminho)
    kwargs = {}
    if senha:
        kwargs["userpw"] = senha
    imagens = pdf2image.convert_from_path(caminho, **kwargs)
    partes = []
    for i, img in enumerate(imagens, start=1):
        try:
            partes.append(pytesseract.image_to_string(img, lang="por"))
        except pytesseract.TesseractError:
            # Sem dados do português instalados, fallback pra inglês
            partes.append(pytesseract.image_to_string(img))
        log.info("OCR página %d: %d chars", i, len(partes[-1]))
    return "\n".join(partes).strip()


# ---------------------------------------------------------------------------
# Pipeline completo (com fallback OCR)
# ---------------------------------------------------------------------------

def pdf_para_linhas(
    caminho: str,
    senha: Optional[str] = None,
) -> tuple[list[dict], dict]:
    """Pipeline completo: PDF -> texto -> parser concreto -> lista de itens.

    Retorna (linhas, metadata) onde:
        linhas: list[dict] no schema {descricao, estabelecimento, categoria,
                                       parcela, valor} — mesmo da pipeline Excel.
        metadata: dict com chaves:
            layout      str   ('nubank', 'itau', 'generico', ...)
            tinha_senha bool  (se o PDF original era protegido)
            n_itens     int   (len(linhas))
            total       float (soma dos valores)
            ocr_usado   bool  (True se foi necessário OCR pra extrair texto)

    Levanta:
        PDFSenhaIncorretaError       — senha ausente/errada
        PDFCorrompidoError           — arquivo inválido
        PDFLayoutDesconhecidoError   — registry vazio (raro)
        PDFCamposIncompletosError    — texto vazio + OCR indisponível
    """
    tinha_senha = pdf_tem_senha(caminho)
    texto = extrair_texto_pdf(caminho, senha=senha)
    ocr_usado = False

    # PDF puramente imagem (escaneado) — texto extraído fica vazio ou minúsculo
    if len(texto) < _LIMITE_TEXTO_VAZIO:
        log.info("Texto extraído curto (%d chars) — tentando OCR", len(texto))
        if not ocr_disponivel():
            raise PDFCamposIncompletosError(
                "Este PDF parece ser escaneado (imagem) e não tem texto extraível. "
                "Para importar, instale o Tesseract OCR e as bibliotecas Python "
                "pytesseract e pdf2image."
            )
        try:
            texto = _extrair_texto_via_ocr(caminho, senha=senha)
            ocr_usado = True
        except Exception as e:
            log.exception("Falha no OCR")
            raise PDFCamposIncompletosError(
                f"Falha ao extrair texto via OCR: {e}"
            ) from e

    from views.cartoes.pdf_parsers import detectar_layout
    parser = detectar_layout(texto)
    log.info("pdf_para_linhas: layout detectado = %s (ocr=%s)",
             parser.nome_layout, ocr_usado)

    linhas = parser.extrair(texto)
    total = round(sum(item.get("valor", 0.0) for item in linhas), 2)

    metadata = {
        "layout":      parser.nome_layout,
        "tinha_senha": tinha_senha,
        "n_itens":     len(linhas),
        "total":       total,
        "ocr_usado":   ocr_usado,
    }
    log.info("pdf_para_linhas: %d itens, total R$ %.2f", len(linhas), total)
    return linhas, metadata


def salvar_como_xlsx(linhas: list[dict], destino: str) -> None:
    """Salva a lista de itens como XLSX no formato compatível com a
    pipeline Excel (mesmas colunas que gerar_template_importacao_xlsx).

    O arquivo gerado pode ser lido de volta por
    importar_fatura_model.ler_arquivo_fatura sem nenhum ajuste.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fatura"

    cabecalhos = ["descricao", "estabelecimento", "categoria", "parcela", "valor"]
    for col, nome in enumerate(cabecalhos, start=1):
        ws.cell(row=1, column=col, value=nome)

    for linha_idx, item in enumerate(linhas, start=2):
        ws.cell(row=linha_idx, column=1, value=item.get("descricao", ""))
        ws.cell(row=linha_idx, column=2, value=item.get("estabelecimento", ""))
        ws.cell(row=linha_idx, column=3, value=item.get("categoria", ""))
        ws.cell(row=linha_idx, column=4, value=item.get("parcela", ""))
        # Valor como número (não string) — Excel reconhece como moeda
        ws.cell(row=linha_idx, column=5, value=float(item.get("valor", 0.0)))

    Path(destino).parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    log.info("salvar_como_xlsx: %d itens em %s", len(linhas), destino)
